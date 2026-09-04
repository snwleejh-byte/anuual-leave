import openpyxl
import datetime
import json
import os
import re

def parse_date(val):
    if not val:
        return None
    if isinstance(val, (datetime.datetime, datetime.date)):
        return datetime.date(val.year, val.month, val.day)
    s = str(val).strip().replace('-', '/').replace('.', '/')
    m = re.match(r'(\d{4})/(\d{1,2})/(\d{1,2})', s)
    if m:
        return datetime.date(int(m.group(1)), int(m.group(2)), int(m.group(3)))
    return None

def normalize_birth(birth_raw, resident_raw):
    # Try parsing birth_raw first
    d = parse_date(birth_raw)
    if d:
        return {
            'date_str': d.strftime('%Y/%m/%d'),
            'birth6': d.strftime('%y%m%d'),
            'birth8': d.strftime('%Y%m%d'),
            'year': d.year,
            'month': d.month,
            'day': d.day
        }
    
    # Try from resident number: e.g. 921108-2000000
    if resident_raw:
        s = str(resident_raw).strip()
        m = re.match(r'(\d{2})(\d{2})(\d{2})-?([1-4])?', s)
        if m:
            yy = int(m.group(1))
            mm = int(m.group(2))
            dd = int(m.group(3))
            gender_code = m.group(4)
            century = 1900
            if gender_code in ['3', '4', '7', '8']:
                century = 2000
            elif gender_code in ['1', '2', '5', '6']:
                century = 1900
            else:
                century = 1900 if yy >= 30 else 2000
            full_year = century + yy
            try:
                real_date = datetime.date(full_year, mm, dd)
                return {
                    'date_str': real_date.strftime('%Y/%m/%d'),
                    'birth6': f"{yy:02d}{mm:02d}{dd:02d}",
                    'birth8': f"{full_year:04d}{mm:02d}{dd:02d}",
                    'year': full_year,
                    'month': mm,
                    'day': dd
                }
            except ValueError:
                pass
    return None

def get_next_anniversary(join_date, ref_date):
    """
    Returns (last_anniversary, next_anniversary, completed_years_at_last_anniv)
    """
    def make_date(year, month, day):
        try:
            return datetime.date(year, month, day)
        except ValueError:
            return datetime.date(year, month, 28)

    candidate = make_date(ref_date.year, join_date.month, join_date.day)
    if candidate <= ref_date:
        last_anniv = candidate
        next_anniv = make_date(ref_date.year + 1, join_date.month, join_date.day)
    else:
        last_anniv = make_date(ref_date.year - 1, join_date.month, join_date.day)
        next_anniv = candidate
        
    completed_years = last_anniv.year - join_date.year
    return last_anniv, next_anniv, completed_years

def calculate_leave(join_date, ref_date=datetime.date(2026, 9, 3)):
    """
    Labor Standards Act Article 60 & SNW Employment Regulation Article 15:
    - 1 year or more:
      - 1st year completed (1주년): 15 days
      - 2nd year completed (2주년): 15 days
      - 3rd year completed (3주년): 16 days
      - 4th year: 16 days
      - 5th year: 17 days
      - formula: 15 + max(0, (completed_years - 1) // 2), max 25 days
    - Less than 1 year:
      - 1 day per completed month, up to max 11 days
      - Valid until the 1-year mark
    """
    total_days_worked = (ref_date - join_date).days
    
    if total_days_worked < 365:
        # Less than 1 year
        m_count = (ref_date.year - join_date.year) * 12 + (ref_date.month - join_date.month)
        if ref_date.day < join_date.day:
            m_count -= 1
        m_count = max(0, min(11, m_count))
        
        try:
            first_anniv = datetime.date(join_date.year + 1, join_date.month, join_date.day)
        except ValueError:
            first_anniv = datetime.date(join_date.year + 1, join_date.month, 28)
            
        return {
            'is_under_1_year': True,
            'completed_years': 0,
            'total_granted': m_count,
            'max_possible': 11,
            'period_start': join_date.strftime('%Y/%m/%d'),
            'period_end': (first_anniv - datetime.timedelta(days=1)).strftime('%Y/%m/%d'),
            'next_renewal_date': first_anniv.strftime('%Y/%m/%d'),
            'rule_description': f'입사 1년 미만 월 단위 발생 ({m_count}개월 개근 / 최대 11일)'
        }
    else:
        last_anniv, next_anniv, completed_years = get_next_anniversary(join_date, ref_date)
        
        if completed_years < 1:
            completed_years = 1
            
        granted = 15
        if completed_years >= 3:
            add_days = (completed_years - 1) // 2
            granted = min(25, 15 + add_days)
            
        period_end = next_anniv - datetime.timedelta(days=1)
        
        return {
            'is_under_1_year': False,
            'completed_years': completed_years,
            'total_granted': granted,
            'max_possible': granted,
            'period_start': last_anniv.strftime('%Y/%m/%d'),
            'period_end': period_end.strftime('%Y/%m/%d'),
            'next_renewal_date': next_anniv.strftime('%Y/%m/%d'),
            'rule_description': f'근속 {completed_years}년차 법정 연차 (기본 15일 + 근속가산 {granted - 15}일)'
        }

def build():
    script_dir = os.path.dirname(os.path.abspath(__file__))
    roster_file = os.path.join(script_dir, '2026.09.03 사원명부.xlsx')
    usage_file = os.path.join(script_dir, '연차 사용 내역 (2025.08.01~2026.08.31).xlsx')
    
    wb_use = openpyxl.load_workbook(usage_file, data_only=True)
    ws_use = wb_use.active
    use_headers = [c for c in next(ws_use.iter_rows(max_row=1, values_only=True))]
    
    u_id_idx = use_headers.index('사번')
    u_name_idx = use_headers.index('성명')
    u_dept_idx = use_headers.index('부서')
    u_date_idx = use_headers.index('근무일')
    u_dow_idx = use_headers.index('요일')
    u_type_idx = use_headers.index('근태')
    u_timecode_idx = use_headers.index('시간코드')
    u_start_idx = use_headers.index('출근')
    u_end_idx = use_headers.index('퇴근')
    u_note_idx = use_headers.index('비고')
    
    usage_by_id = {}
    usage_by_name = {}
    
    total_usage_count = 0
    for r in ws_use.iter_rows(min_row=2, values_only=True):
        if not r[u_name_idx]:
            continue
        emp_id = str(r[u_id_idx]).strip() if r[u_id_idx] else ''
        name = str(r[u_name_idx]).strip()
        
        g_type = str(r[u_type_idx]).strip()
        deduct = 1.0 if '년차' in g_type else 0.5
        
        date_str = str(r[u_date_idx]).strip()
        d_parsed = parse_date(date_str)
        date_formatted = d_parsed.strftime('%Y/%m/%d') if d_parsed else date_str
        
        item = {
            'emp_id': emp_id,
            'name': name,
            'dept': str(r[u_dept_idx]).strip() if r[u_dept_idx] else '',
            'date': date_formatted,
            'date_obj': d_parsed,
            'day_of_week': str(r[u_dow_idx]).strip() if r[u_dow_idx] else '',
            'leave_type': g_type,
            'days': deduct,
            'time_code': str(r[u_timecode_idx]).strip() if r[u_timecode_idx] else '',
            'start_time': str(r[u_start_idx]).strip() if r[u_start_idx] else '',
            'end_time': str(r[u_end_idx]).strip() if r[u_end_idx] else '',
            'note': str(r[u_note_idx]).strip() if r[u_note_idx] else ''
        }
        
        if emp_id:
            usage_by_id.setdefault(emp_id, []).append(item)
        usage_by_name.setdefault(name, []).append(item)
        total_usage_count += 1
        
    print(f"Loaded {total_usage_count} usage rows.")
    
    wb_emp = openpyxl.load_workbook(roster_file, data_only=True)
    ws_emp = wb_emp.active
    emp_headers = [c for c in next(ws_emp.iter_rows(max_row=1, values_only=True))]
    
    e_workplace_idx = emp_headers.index('사업장')
    e_dept_idx = emp_headers.index('부서')
    e_id_idx = emp_headers.index('사번')
    e_name_idx = emp_headers.index('성명')
    e_pos_idx = emp_headers.index('직위')
    e_rank_idx = emp_headers.index('직급')
    e_job_idx = emp_headers.index('직책') if '직책' in emp_headers else -1
    e_duty_idx = emp_headers.index('직무') if '직무' in emp_headers else -1
    e_resident_idx = emp_headers.index('주민등록번호')
    e_emp_type_idx = emp_headers.index('사원구분')
    e_gender_idx = emp_headers.index('성별')
    e_birth_idx = emp_headers.index('생년월일')
    e_join_idx = emp_headers.index('입사일자')
    e_retire_idx = emp_headers.index('퇴사일자')
    e_phone_idx = emp_headers.index('비상연락') if '비상연락' in emp_headers else -1
    e_email_idx = emp_headers.index('이메일') if '이메일' in emp_headers else -1
    e_shift_idx = emp_headers.index('근무조') if '근무조' in emp_headers else -1
    
    employees = []
    ref_date = datetime.date(2026, 9, 3)
    
    for r in ws_emp.iter_rows(min_row=2, values_only=True):
        if not r[e_name_idx]:
            continue
        name = str(r[e_name_idx]).strip()
        emp_id = str(r[e_id_idx]).strip()
        dept = str(r[e_dept_idx]).strip() if r[e_dept_idx] else '-'
        position = str(r[e_pos_idx]).strip() if r[e_pos_idx] else ''
        rank = str(r[e_rank_idx]).strip() if r[e_rank_idx] else ''
        emp_type = str(r[e_emp_type_idx]).strip() if r[e_emp_type_idx] else ''
        gender = str(r[e_gender_idx]).strip() if r[e_gender_idx] else ''
        shift = str(r[e_shift_idx]).strip() if (e_shift_idx != -1 and r[e_shift_idx]) else ''
        email = str(r[e_email_idx]).strip() if (e_email_idx != -1 and r[e_email_idx]) else ''
        phone = str(r[e_phone_idx]).strip() if (e_phone_idx != -1 and r[e_phone_idx]) else ''
        retire = str(r[e_retire_idx]).strip() if r[e_retire_idx] else None
        
        join_date = parse_date(r[e_join_idx])
        if not join_date:
            print(f"Warning: Missing join date for {name} ({emp_id})")
            continue
            
        birth_info = normalize_birth(r[e_birth_idx], r[e_resident_idx])
        if not birth_info:
            print(f"Warning: Missing birth info for {name} ({emp_id})")
            continue
            
        days_worked = (ref_date - join_date).days
        years_worked = days_worked // 365
        rem_days = days_worked % 365
        months_approx = rem_days // 30
        
        calc = calculate_leave(join_date, ref_date)
        
        raw_usage = usage_by_id.get(emp_id) or usage_by_name.get(name) or []
        
        p_start = parse_date(calc['period_start'])
        p_end = parse_date(calc['period_end'])
        
        current_period_usage = []
        prior_usage = []
        
        total_used_current = 0.0
        for u in raw_usage:
            clean_item = {
                'date': u['date'],
                'day_of_week': u['day_of_week'],
                'leave_type': u['leave_type'],
                'days': u['days'],
                'time_code': u['time_code'],
                'start_time': u['start_time'],
                'end_time': u['end_time'],
                'note': u['note']
            }
            if u['date_obj'] and p_start and p_end:
                if p_start <= u['date_obj'] <= p_end:
                    current_period_usage.append(clean_item)
                    total_used_current += u['days']
                else:
                    prior_usage.append(clean_item)
            else:
                current_period_usage.append(clean_item)
                total_used_current += u['days']
                
        current_period_usage.sort(key=lambda x: x['date'], reverse=True)
        prior_usage.sort(key=lambda x: x['date'], reverse=True)
        all_usage = current_period_usage + prior_usage
        all_usage.sort(key=lambda x: x['date'], reverse=True)
        
        remaining_days = round(calc['total_granted'] - total_used_current, 2)
        usage_rate = round((total_used_current / calc['total_granted'] * 100) if calc['total_granted'] > 0 else 0, 1)
        
        renewal_date = parse_date(calc['next_renewal_date'])
        d_day = (renewal_date - ref_date).days if renewal_date else 0
        
        emp_obj = {
            'emp_id': emp_id,
            'name': name,
            'dept': dept,
            'position': position,
            'rank': rank,
            'emp_type': emp_type,
            'gender': gender,
            'shift': shift,
            'email': email,
            'phone': phone,
            'retire_date': retire,
            'birth_info': birth_info,
            'join_date': join_date.strftime('%Y/%m/%d'),
            'service_years': years_worked,
            'service_months': months_approx,
            'service_days': days_worked,
            'service_text': f"{years_worked}년 {months_approx}개월 ({days_worked}일)" if years_worked > 0 else f"{months_approx}개월 ({days_worked}일)",
            'leave_calc': {
                'is_under_1_year': calc['is_under_1_year'],
                'completed_years': calc['completed_years'],
                'total_granted': calc['total_granted'],
                'max_possible': calc['max_possible'],
                'used_days': total_used_current,
                'remaining_days': remaining_days,
                'usage_rate': usage_rate,
                'period_start': calc['period_start'],
                'period_end': calc['period_end'],
                'next_renewal_date': calc['next_renewal_date'],
                'd_day': d_day,
                'rule_description': calc['rule_description']
            },
            'current_usage': current_period_usage,
            'all_usage': all_usage
        }
        employees.append(emp_obj)
        
    print(f"Processed {len(employees)} active employees.")
    
    employees.sort(key=lambda x: x['name'])
    
    total_granted_all = sum(e['leave_calc']['total_granted'] for e in employees)
    total_used_all = sum(e['leave_calc']['used_days'] for e in employees)
    total_rem_all = sum(e['leave_calc']['remaining_days'] for e in employees)
    
    dataset = {
        'company': '(주)에스앤더블류',
        'ref_date': ref_date.strftime('%Y/%m/%d'),
        'total_employees': len(employees),
        'summary': {
            'total_granted': round(total_granted_all, 1),
            'total_used': round(total_used_all, 1),
            'total_remaining': round(total_rem_all, 1),
            'avg_usage_rate': round(total_used_all / total_granted_all * 100, 1) if total_granted_all > 0 else 0
        },
        'employees': employees
    }
    
    json_path = os.path.join(script_dir, 'data.json')
    js_path = os.path.join(script_dir, 'data.js')
    
    with open(json_path, 'w', encoding='utf-8') as f:
        json.dump(dataset, f, ensure_ascii=False, indent=2)
        
    with open(js_path, 'w', encoding='utf-8') as f:
        f.write('// Auto-generated SNW Leave Dataset\n')
        f.write('window.SNW_DATA = ')
        json.dump(dataset, f, ensure_ascii=False, indent=2)
        f.write(';\n')
        
    print("Successfully generated data.json and data.js!")

if __name__ == '__main__':
    build()
